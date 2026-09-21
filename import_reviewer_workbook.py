#!/usr/bin/env python3
"""Validate a completed reviewer workbook and export canonical decision JSONL."""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


HEADERS = [
    "record_key", "title", "abstract", "doi", "year", "landing_page_url",
    "metadata_problem", "container_title", "container_resource_url", "decision",
    "reason", "decided_at_utc", "source_locator",
]
DECISIONS = {"retrieve_full_text", "exclude", "background_method"}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text().splitlines() if line]


def normalize_timestamp(value: Any) -> str:
    if isinstance(value, datetime):
        parsed = value.replace(tzinfo=value.tzinfo or timezone.utc)
    elif isinstance(value, str) and value.strip():
        parsed = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
        parsed = parsed.replace(tzinfo=parsed.tzinfo or timezone.utc)
    else:
        raise ValueError("missing decided_at_utc")
    return parsed.astimezone(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--packet", required=True, type=Path)
    parser.add_argument("--attachment", required=True, type=Path)
    parser.add_argument("--reviewer", required=True, choices=["A", "B"])
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()

    packet = read_jsonl(args.packet)
    attachment = {row["record_key"]: row for row in read_jsonl(args.attachment)}
    workbook = load_workbook(args.workbook, read_only=True, data_only=False)
    if workbook.sheetnames != ["Summary", "Review", "Instructions"]:
        raise ValueError(f"unexpected workbook sheets: {workbook.sheetnames}")
    rows = workbook["Review"].iter_rows(values_only=True)
    header = list(next(rows))
    if header != HEADERS:
        raise ValueError("review sheet headers changed")

    decisions: list[dict[str, Any]] = []
    errors: list[str] = []
    for line_number, source in enumerate(packet, 2):
        row = next(rows, None)
        if row is None:
            errors.append(f"row {line_number}: missing workbook row")
            continue
        extra = attachment.get(source["record_key"], {})
        expected = (
            source["record_key"], source.get("title"), source.get("abstract"), source.get("doi"),
            source.get("year"), source.get("landing_page_url"), source.get("metadata_problem"),
            extra.get("container_title"), extra.get("primary_resource_url"),
        )
        if tuple(row[:9]) != expected:
            errors.append(f"row {line_number}: source metadata changed for {source['record_key']}")
            continue
        decision, reason, decided_at, locator = row[9:13]
        if decision not in DECISIONS:
            errors.append(f"row {line_number}: missing or invalid decision for {source['record_key']}")
            continue
        if not isinstance(reason, str) or not reason.strip():
            errors.append(f"row {line_number}: missing reason for {source['record_key']}")
            continue
        if not isinstance(locator, str) or not locator.strip():
            errors.append(f"row {line_number}: missing source locator for {source['record_key']}")
            continue
        try:
            timestamp = normalize_timestamp(decided_at)
        except (ValueError, TypeError) as exc:
            errors.append(f"row {line_number}: {exc} for {source['record_key']}")
            continue
        decisions.append({
            "record_key": source["record_key"], "reviewer_id": args.reviewer,
            "stage": "title_abstract", "decision": decision, "reason": reason.strip(),
            "decided_at_utc": timestamp, "source_locator": locator.strip(),
        })
    if next(rows, None) is not None:
        errors.append("workbook contains extra review rows")
    if errors:
        preview = "\n".join(errors[:20])
        suffix = f"\n... {len(errors) - 20} more" if len(errors) > 20 else ""
        raise ValueError(f"workbook rejected with {len(errors)} error(s):\n{preview}{suffix}")
    if len(decisions) != len(packet):
        raise ValueError("decision count does not match canonical packet")

    body = "".join(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n" for row in decisions)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    if args.out.exists() and args.out.read_text() != body:
        raise ValueError(f"refusing to overwrite differing decision file: {args.out}")
    args.out.write_text(body)
    manifest_path = args.out.with_suffix(args.out.suffix + ".manifest.json")
    manifest = {
        "status": "validated_complete_reviewer_decisions", "reviewer_id": args.reviewer,
        "workbook_file": args.workbook.name, "workbook_sha256": sha256(args.workbook),
        "packet_file": args.packet.name, "packet_sha256": sha256(args.packet),
        "attachment_sha256": sha256(args.attachment), "decision_records": len(decisions),
        "decision_file": args.out.name, "decision_sha256": hashlib.sha256(body.encode()).hexdigest(),
    }
    manifest_text = json.dumps(manifest, indent=2) + "\n"
    if manifest_path.exists() and manifest_path.read_text() != manifest_text:
        raise ValueError(f"refusing to overwrite differing manifest: {manifest_path}")
    manifest_path.write_text(manifest_text)
    print(json.dumps(manifest, indent=2))


if __name__ == "__main__":
    main()
