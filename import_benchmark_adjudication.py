#!/usr/bin/env python3
"""Validate a completed independent benchmark workbook and export decisions."""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


HEADERS = [
    "row_key", "case_id", "claim", "dataset_id", "platform", "collection_period",
    "unit", "claim_subject", "evidence", "source", "source_locator", "source_sha256",
    "source_hash_scope", "reviewer_label", "rationale", "source_checked", "decided_at_utc",
]
SOURCE_FIELDS = HEADERS[:13]
LABELS = {"supported", "contradicted", "unknown"}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text().splitlines() if line]


def normalize_timestamp(value: Any, field: str) -> str:
    if isinstance(value, datetime):
        parsed = value.replace(tzinfo=value.tzinfo or timezone.utc)
    elif isinstance(value, str) and value.strip():
        parsed = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
        parsed = parsed.replace(tzinfo=parsed.tzinfo or timezone.utc)
    else:
        raise ValueError(f"missing {field}")
    return parsed.astimezone(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--packet", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()

    packet = read_jsonl(args.packet)
    workbook = load_workbook(args.workbook, read_only=True, data_only=False)
    if workbook.sheetnames != ["Summary", "Adjudication", "Instructions"]:
        raise ValueError(f"unexpected workbook sheets: {workbook.sheetnames}")

    summary = workbook["Summary"]
    reviewer_id = summary["B11"].value
    confirmation = summary["B12"].value
    conflict = summary["B13"].value
    completed_at = summary["B14"].value
    if not isinstance(reviewer_id, str) or not reviewer_id.strip():
        raise ValueError("missing reviewer identifier in Summary!B11")
    if confirmation != "confirmed":
        raise ValueError("independent review is not confirmed in Summary!B12")
    if not isinstance(conflict, str) or not conflict.strip():
        raise ValueError("missing conflict disclosure in Summary!B13; enter 'none' when applicable")
    completed_at_utc = normalize_timestamp(completed_at, "review completed timestamp in Summary!B14")

    rows = workbook["Adjudication"].iter_rows(values_only=True)
    if list(next(rows)) != HEADERS:
        raise ValueError("adjudication headers changed")

    decisions: list[dict[str, Any]] = []
    errors: list[str] = []
    for line_number, source in enumerate(packet, 2):
        row = next(rows, None)
        if row is None:
            errors.append(f"row {line_number}: missing workbook row")
            continue
        expected = tuple(source[field] for field in SOURCE_FIELDS)
        if tuple(row[:13]) != expected:
            errors.append(f"row {line_number}: source metadata changed for {source['row_key']}")
            continue
        label, rationale, checked, decided_at = row[13:17]
        if label not in LABELS:
            errors.append(f"row {line_number}: missing or invalid label for {source['row_key']}")
            continue
        if not isinstance(rationale, str) or not rationale.strip():
            errors.append(f"row {line_number}: missing rationale for {source['row_key']}")
            continue
        if checked != "yes":
            errors.append(f"row {line_number}: source not confirmed as checked for {source['row_key']}")
            continue
        try:
            timestamp = normalize_timestamp(decided_at, "decision timestamp")
        except (TypeError, ValueError) as exc:
            errors.append(f"row {line_number}: {exc} for {source['row_key']}")
            continue
        decisions.append({
            "row_key": source["row_key"],
            "case_id": source["case_id"],
            "claim": source["claim"],
            "reviewer_id": reviewer_id.strip(),
            "label": label,
            "rationale": rationale.strip(),
            "source_checked": True,
            "decided_at_utc": timestamp,
        })
    if next(rows, None) is not None:
        errors.append("workbook contains extra adjudication rows")
    if errors:
        preview = "\n".join(errors[:20])
        suffix = f"\n... {len(errors) - 20} more" if len(errors) > 20 else ""
        raise ValueError(f"workbook rejected with {len(errors)} error(s):\n{preview}{suffix}")
    if len(decisions) != len(packet):
        raise ValueError("decision count does not match canonical packet")

    body = "".join(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n" for row in decisions)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    if args.out.exists() and args.out.read_text() != body:
        raise ValueError(f"refusing to overwrite differing decision file: {args.out}")
    args.out.write_text(body)
    manifest_path = args.out.with_suffix(args.out.suffix + ".manifest.json")
    manifest = {
        "status": "validated_complete_independent_benchmark_adjudication",
        "reviewer_id": reviewer_id.strip(),
        "independent_review_confirmation": confirmation,
        "conflict_disclosure": conflict.strip(),
        "review_completed_at_utc": completed_at_utc,
        "workbook_file": args.workbook.name,
        "workbook_sha256": sha256(args.workbook),
        "packet_file": args.packet.name,
        "packet_sha256": sha256(args.packet),
        "decision_records": len(decisions),
        "decision_file": args.out.name,
        "decision_sha256": hashlib.sha256(body.encode()).hexdigest(),
    }
    manifest_text = json.dumps(manifest, indent=2) + "\n"
    if manifest_path.exists() and manifest_path.read_text() != manifest_text:
        raise ValueError(f"refusing to overwrite differing manifest: {manifest_path}")
    manifest_path.write_text(manifest_text)
    print(manifest_text, end="")


if __name__ == "__main__":
    main()
