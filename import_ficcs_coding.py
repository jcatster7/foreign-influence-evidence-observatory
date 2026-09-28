#!/usr/bin/env python3
import argparse, hashlib, json
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
PACKET = ROOT / 'ficcs/coding/blind_claim_packet_v1.jsonl'
CONSTRUCTS = tuple('OPACDERI')
EDGES = (('A','O'),('C','P'),('P','A'),('C','D'),('engagement','E'),('E','R'),('R','I'))
CONSEQUENCES = ('none_descriptive','account_enforcement','content_enforcement','public_attribution','sanctions_or_diplomacy','legislation_or_regulation','resource_allocation','public_warning','other')

def rows(ws):
    values = list(ws.iter_rows(values_only=True))
    head = [str(x) for x in values[0]]
    return [dict(zip(head, row)) for row in values[1:]]

def require(value, label, allowed=None):
    assert value not in (None, ''), f'missing {label}'
    if allowed: assert value in allowed, f'invalid {label}: {value}'
    return value

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('workbook')
    ap.add_argument('--out', required=True)
    args = ap.parse_args()
    claims = [json.loads(x) for x in PACKET.read_text().splitlines() if x]
    by_id = {x['claim_id']: x for x in claims}
    wb = load_workbook(args.workbook, read_only=True, data_only=False)
    assert set(('Summary','Constructs','Edges','Consequences')).issubset(wb.sheetnames)
    summary = {r[0].value:r[1].value for r in wb['Summary'].iter_rows(min_row=2, max_col=2)}
    role = require(summary['Reviewer role'], 'reviewer role', ('A','B'))
    reviewer = str(require(summary['Reviewer ID'], 'reviewer ID')).strip()
    conflict = str(require(summary['Conflict disclosure'], 'conflict disclosure')).strip()
    require(summary['No-collaboration attestation'], 'attestation', ('I attest',))
    completed = str(require(summary['Completed at UTC'], 'completion timestamp')).strip()
    datetime.fromisoformat(completed.replace('Z','+00:00'))

    construct_rows = rows(wb['Constructs'])
    edge_rows = rows(wb['Edges'])
    consequence_rows = rows(wb['Consequences'])
    assert len(construct_rows) == len(claims)*8
    assert len(edge_rows) == len(claims)*7
    assert len(consequence_rows) == len(claims)
    decisions = []
    for i, row in enumerate(construct_rows):
        claim = claims[i//8]; code = CONSTRUCTS[i%8]
        assert row['claim_id'] == claim['claim_id'] and row['document_id'] == claim['source']['document_id'] and row['construct'] == code
        assert row['claim_paraphrase'] == claim['claim']['paraphrase'] and row['canonical_url'] == claim['source']['canonical_url']
        decisions.append({'row_type':'construct','row_key':f"{claim['claim_id']}:{code}",'claim_id':claim['claim_id'],'construct':code,'asserted_by_source':row['asserted_by_source']=='yes','decision':require(row['decision'],'construct decision',('supported','contradicted','unknown')),'evidence_stage':require(row['evidence_stage'],'evidence stage',('direct','proxy','platform_disclosure','measurement_validation','unknown')),'evidence_locator':str(require(row['evidence_locator'],'evidence locator')).strip(),'rationale':str(require(row['rationale'],'construct rationale')).strip(),'confidence':require(row['confidence'],'confidence',('high','moderate','low')),'source_checked':require(row['source_checked'],'source checked',('yes',))=='yes','decided_at_utc':str(require(row['decided_at_utc'],'decision timestamp')).strip()})
    for i, row in enumerate(edge_rows):
        claim = claims[i//7]; edge = EDGES[i%7]
        assert row['claim_id'] == claim['claim_id'] and (row['from'],row['to']) == edge and row['claim_paraphrase'] == claim['claim']['paraphrase']
        decisions.append({'row_type':'edge','row_key':f"{claim['claim_id']}:{edge[0]}>{edge[1]}",'claim_id':claim['claim_id'],'from':edge[0],'to':edge[1],'asserted':require(row['asserted'],'edge asserted',('yes','no','unclear')),'status':require(row['status'],'edge status',('supported_link','measured_association','proxy_only','untested','contradicted_link')),'rationale':str(require(row['rationale'],'edge rationale')).strip(),'source_checked':require(row['source_checked'],'source checked',('yes',))=='yes','decided_at_utc':str(require(row['decided_at_utc'],'decision timestamp')).strip()})
    for row in consequence_rows:
        claim = by_id.get(row['claim_id']); assert claim and row['document_id'] == claim['source']['document_id'] and row['claim_paraphrase'] == claim['claim']['paraphrase']
        categories = [c for c in CONSEQUENCES if require(row[c],c,('yes','no')) == 'yes']
        assert categories, f"no consequence category: {claim['claim_id']}"
        decisions.append({'row_type':'consequence','row_key':f"{claim['claim_id']}:consequence",'claim_id':claim['claim_id'],'categories':categories,'action_basis':require(row['action_basis'],'action basis',('none_descriptive','recommended','documented_taken','both','unclear')),'severity':require(row['severity'],'severity',('low','moderate','high')),'rationale':str(require(row['rationale'],'consequence rationale')).strip(),'source_checked':require(row['source_checked'],'source checked',('yes',))=='yes','decided_at_utc':str(require(row['decided_at_utc'],'decision timestamp')).strip()})
    out = Path(args.out); out.parent.mkdir(parents=True, exist_ok=True)
    text = ''.join(json.dumps(x,separators=(',',':'))+'\n' for x in decisions); out.write_text(text)
    source_bytes = Path(args.workbook).read_bytes()
    manifest = {'status':'validated_complete_independent_ficcs_coding','reviewer_role':role,'reviewer_id':reviewer,'conflict_disclosure':conflict,'workbook_sha256':hashlib.sha256(source_bytes).hexdigest(),'packet_sha256':hashlib.sha256(PACKET.read_bytes()).hexdigest(),'decision_sha256':hashlib.sha256(text.encode()).hexdigest(),'rows':len(decisions),'construct_rows':len(construct_rows),'edge_rows':len(edge_rows),'consequence_rows':len(consequence_rows),'completed_at_utc':completed}
    Path(str(out)+'.manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
    print(json.dumps(manifest,indent=2))

if __name__ == '__main__': main()
